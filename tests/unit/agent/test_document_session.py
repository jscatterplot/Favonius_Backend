"""Behavioural tests for controller._run_document_fill_turn.

Isolates the DB + LLM seams (document_store, budget, run_qa_turn, audit
writers, the Anthropic client getter) via monkeypatch, but uses the REAL
``document_render`` against a real in-memory DOCX so the grounding-drop +
render + reply-shape logic is exercised end to end. Covers:

- finalize: valid field filled, UNKNOWN field + UNANCHORED replacement dropped,
  output stored as 'final', session → 'finalized', reply carries a download.
- ask: reply.status='needs_input' with questions, session → 'awaiting_input',
  nothing rendered (no values yet).
"""

from __future__ import annotations

import io
from typing import Any
from uuid import uuid4

import pytest

import src.api.agent.budget as budget_mod
import src.api.agent.controller as controller
import src.api.agent.document_store as ds
import src.api.agent.llm as agent_llm
from src.api.agent.auth_context import AuthContext
from src.api.agent.controller import _run_document_fill_turn
from src.api.agent_workflows.models import ToolCall
from src.api.agent_workflows.runtime import EMIT_FINAL_ANSWER_TOOL, QAResult


def _template_docx() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Report for {{ total }}.")
    doc.add_paragraph("Period: April 2026.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _auth() -> AuthContext:
    return AuthContext(
        user_id=uuid4(),
        organization_id=uuid4(),
        role="customer_admin",
        visible_depot_ids=[uuid4()],
    )


@pytest.fixture
def patched(monkeypatch):
    """Patch the DB/LLM seams; return a dict of recorders for assertions."""
    template_id = uuid4()
    session_id = uuid4()
    state: dict[str, Any] = {
        "session_id": session_id,
        "template_id": template_id,
        "store_output_calls": [],
        "update_session_calls": [],
        "session_status": "gathering",
        "session_draft": controller.initial_draft(),
        "message_log": [],
        "qa": None,  # set per-test
    }

    async def _load_session(ts_pool, sid, user_id, *, is_admin):
        return {
            "id": session_id,
            "template_id": template_id,
            "organization_id": uuid4(),
            "depot_id": None,
            "user_id": user_id,
            "status": state["session_status"],
            "draft": state["session_draft"],
            "message_log": state["message_log"],
            "latest_output_id": None,
        }

    async def _load_template(ts_pool, tid):
        return {
            "id": template_id,
            "organization_id": uuid4(),
            "depot_id": None,
            "kind": "docx",
            "pdf_form_type": None,
            "file_name": "report.docx",
            "raw_payload": _template_docx(),
            "detected_fields": [{"name": "total", "source": "jinja"}],
            "status": "stored",
        }

    async def _store_output(ts_pool, **kwargs):
        state["store_output_calls"].append(kwargs)
        return uuid4()

    async def _update_session(ts_pool, sid, **kwargs):
        state["update_session_calls"].append(kwargs)

    monkeypatch.setattr(ds, "load_session_for_user", _load_session)
    monkeypatch.setattr(ds, "load_template", _load_template)
    monkeypatch.setattr(ds, "store_output", _store_output)
    monkeypatch.setattr(ds, "update_session", _update_session)

    # Budget — always allow; reconcile no-op.
    async def _check_and_reserve(org_id, est):
        return object()

    monkeypatch.setattr(budget_mod, "check_and_reserve", _check_and_reserve)
    monkeypatch.setattr(budget_mod, "record_actual", lambda *a, **k: None)

    # Audit writers — no-op (ts_pool is None in the test).
    async def _noop(*a, **k):
        return None

    monkeypatch.setattr(controller, "agent_runs_step", _noop)
    monkeypatch.setattr(controller, "agent_runs_close", _noop)
    monkeypatch.setattr(controller, "write_agent_query_audit", _noop)

    # Anthropic client getter — never actually used (run_qa_turn is faked).
    monkeypatch.setattr(agent_llm, "_get_client", lambda: object())

    # run_qa_turn — return whatever QAResult the test set on state["qa"].
    async def _fake_run_qa_turn(**kwargs):
        return state["qa"]

    monkeypatch.setattr(controller, "run_qa_turn", _fake_run_qa_turn)

    return state


def _terminator_qa(args: dict[str, Any]) -> QAResult:
    tc = ToolCall(
        name=EMIT_FINAL_ANSWER_TOOL,
        arguments=args,
        result={"text": args.get("text", ""), "row_evidence": 0, "mode": args.get("mode")},
        ok=True,
    )
    return QAResult(
        text=args.get("text", ""),
        tool_calls=[tc],
        status="success",
        row_evidence=0,
        iterations=2,
        input_tokens=10,
        output_tokens=5,
    )


async def _emit_step(name, *, label_key=None):  # noqa: D401 - test noop
    return None


@pytest.mark.asyncio
async def test_finalize_fills_valid_drops_ungrounded_and_renders(patched):
    patched["qa"] = _terminator_qa(
        {
            "mode": "finalize",
            "text": "I've updated the report.",
            "field_values": {"total": "15,000 kWh", "bogus_field": "x"},
            "replacements": [
                {"find": "April 2026", "replace": "May 2026"},
                {"find": "NONEXISTENT TEXT", "replace": "z"},
            ],
            "questions": [],
        }
    )
    reply = await _run_document_fill_turn(
        run_id=uuid4(),
        message="finalize it for May 2026",
        auth=_auth(),
        static_pool=None,
        ts_pool=None,
        session_id=patched["session_id"],
        sse=None,
        emit_step=_emit_step,
    )

    assert reply.status == "success"
    assert reply.intent == "document_fill"
    assert reply.session_id == patched["session_id"]
    assert reply.download is not None
    assert reply.download["output_kind"] == "final"
    assert reply.download["kind"] == "docx"
    assert reply.download["fidelity"] == "preserved"

    # Grounding: the stored output got ONLY the known field + anchored replacement.
    assert len(patched["store_output_calls"]) == 1
    stored = patched["store_output_calls"][0]
    assert stored["output_kind"] == "final"
    assert "total" in stored["field_values"]
    assert "bogus_field" not in stored["field_values"]
    assert [r["find"] for r in stored["replacements"]] == ["April 2026"]

    # The dropped items are surfaced to the user.
    assert "bogus_field" in reply.text
    assert "didn't match" in reply.text

    # Session advanced to finalized.
    assert len(patched["update_session_calls"]) == 1
    upd = patched["update_session_calls"][0]
    assert upd["status"] == "finalized"
    assert upd["draft"]["field_values"]["total"] == "15,000 kWh"


@pytest.mark.asyncio
async def test_ask_pauses_with_questions_and_no_render(patched):
    patched["qa"] = _terminator_qa(
        {
            "mode": "ask",
            "text": "Which period should this cover?",
            "field_values": {},
            "replacements": [],
            "questions": [{"id": "q1", "text": "Which period should this cover?"}],
        }
    )
    reply = await _run_document_fill_turn(
        run_id=uuid4(),
        message="update this report",
        auth=_auth(),
        static_pool=None,
        ts_pool=None,
        session_id=patched["session_id"],
        sse=None,
        emit_step=_emit_step,
    )

    assert reply.status == "needs_input"
    assert reply.session_id == patched["session_id"]
    assert reply.questions == [{"id": "q1", "text": "Which period should this cover?"}]
    # Nothing to render yet → no output stored, no download.
    assert patched["store_output_calls"] == []
    assert reply.download is None
    # Session paused awaiting the user.
    assert patched["update_session_calls"][0]["status"] == "awaiting_input"


@pytest.mark.asyncio
async def test_unknown_session_returns_not_found(patched, monkeypatch):
    async def _no_session(*a, **k):
        return None

    monkeypatch.setattr(ds, "load_session_for_user", _no_session)
    reply = await _run_document_fill_turn(
        run_id=uuid4(),
        message="hi",
        auth=_auth(),
        static_pool=None,
        ts_pool=None,
        session_id=uuid4(),
        sse=None,
        emit_step=_emit_step,
    )
    assert reply.status == "not_found"
    assert reply.intent == "document_fill"
