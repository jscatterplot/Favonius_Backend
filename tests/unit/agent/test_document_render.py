"""Unit tests for src/api/agent/document_render.py (pure, no DB/LLM).

Covers extraction + rendering round-trips for all three paths:
- DOCX: jinja field fill (docxtpl) + targeted old→new replacement, layout kept.
- Flat PDF: text extraction + degraded re-render.
- AcroForm PDF: field detection + fill, layout kept.

PDF tests importorskip pypdf/reportlab so a DOCX-only environment still runs the
DOCX cases.
"""

from __future__ import annotations

import io

import pytest

from src.api.agent import document_render as dr


def _docx_bytes(paragraphs: list[str]) -> bytes:
    from docx import Document

    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _docx_text(body: bytes) -> str:
    return dr.extract(body, "docx").text


# ── sniff_kind ───────────────────────────────────────────────────────────────


def test_sniff_kind_docx():
    body = _docx_bytes(["hello"])
    assert dr.sniff_kind(body) == "docx"


def test_sniff_kind_pdf():
    assert dr.sniff_kind(b"%PDF-1.7\n...") == "pdf"


def test_sniff_kind_unknown():
    assert dr.sniff_kind(b"just plain text, not a document") is None
    assert dr.sniff_kind(b"") is None


# ── DOCX extraction + fill ─────────────────────────────────────────────────


def test_extract_docx_detects_jinja_fields():
    body = _docx_bytes(["Report for {{ customer_name }} covering {{ period }}."])
    ex = dr.extract(body, "docx")
    names = {f.name for f in ex.fields}
    assert names == {"customer_name", "period"}
    assert all(f.source == "jinja" for f in ex.fields)
    assert ex.pdf_form_type is None


def test_extract_docx_old_report_has_no_fields():
    # A finished report (no placeholders) → empty fields; agent uses replacements.
    body = _docx_bytes(["Total energy consumed: 12,340 kWh in April 2026."])
    ex = dr.extract(body, "docx")
    assert ex.fields == []
    assert "12,340 kWh" in ex.text


def test_render_docx_fills_jinja_and_applies_replacement():
    body = _docx_bytes(
        [
            "Monthly report for {{ customer_name }}.",
            "Total energy consumed: 12,340 kWh in April 2026.",
        ]
    )
    out, fidelity = dr.render(
        kind="docx",
        pdf_form_type=None,
        body=body,
        field_values={"customer_name": "ACME Fleet"},
        replacements=[{"find": "12,340 kWh", "replace": "15,000 kWh"}],
    )
    assert fidelity == "preserved"
    text = _docx_text(out)
    assert "ACME Fleet" in text
    assert "15,000 kWh" in text
    assert "12,340 kWh" not in text
    assert "{{" not in text  # tag consumed


def test_render_docx_replacement_only_old_report():
    body = _docx_bytes(["Period: April 2026. Total: 12,340 kWh."])
    out, fidelity = dr.render(
        kind="docx",
        pdf_form_type=None,
        body=body,
        field_values={},
        replacements=[
            {"find": "April 2026", "replace": "May 2026"},
            {"find": "12,340 kWh", "replace": "15,000 kWh"},
        ],
    )
    assert fidelity == "preserved"
    text = _docx_text(out)
    assert "May 2026" in text and "15,000 kWh" in text
    assert "April 2026" not in text and "12,340" not in text


def test_render_docx_split_run_anchor_falls_back_to_paragraph_join():
    # Build a paragraph whose target text is split across two runs — the
    # per-run fast path misses it, so the paragraph-join fallback must catch it.
    from docx import Document

    doc = Document()
    para = doc.add_paragraph()
    para.add_run("Total: 12,")
    para.add_run("340 kWh")
    buf = io.BytesIO()
    doc.save(buf)
    body = buf.getvalue()
    assert "12,340 kWh" in _docx_text(body)

    out, _ = dr.render(
        kind="docx",
        pdf_form_type=None,
        body=body,
        field_values={},
        replacements=[{"find": "12,340 kWh", "replace": "15,000 kWh"}],
    )
    text = _docx_text(out)
    assert "15,000 kWh" in text
    assert "12,340" not in text


def test_extract_unknown_kind_raises():
    with pytest.raises(dr.DocumentRenderError):
        dr.extract(b"x", "rtf")


# ── Flat PDF ────────────────────────────────────────────────────────────────


def _flat_pdf(text: str) -> bytes:
    canvas = pytest.importorskip("reportlab.pdfgen.canvas")
    from reportlab.lib.pagesizes import A4

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.drawString(72, 720, text)
    c.save()
    return buf.getvalue()


def test_extract_flat_pdf():
    pytest.importorskip("pypdf")
    body = _flat_pdf("Depot report. Total: 12340 kWh.")
    ex = dr.extract(body, "pdf")
    assert ex.pdf_form_type == "flat"
    assert ex.fields == []
    assert "12340" in ex.text


def test_render_flat_pdf_is_degraded_and_applies_replacement():
    pytest.importorskip("pypdf")
    body = _flat_pdf("Depot report. Total: 12340 kWh.")
    ex = dr.extract(body, "pdf")
    out, fidelity = dr.render(
        kind="pdf",
        pdf_form_type="flat",
        body=body,
        field_values={},
        replacements=[{"find": "12340", "replace": "15000"}],
        full_text=ex.text,
    )
    assert fidelity == "degraded"
    assert "15000" in dr.extract(out, "pdf").text


# ── AcroForm PDF ────────────────────────────────────────────────────────────


def _acroform_pdf() -> bytes:
    canvas = pytest.importorskip("reportlab.pdfgen.canvas")
    from reportlab.lib.pagesizes import A4

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.drawString(72, 740, "Customer:")
    c.acroForm.textfield(
        name="customer_name", x=160, y=735, width=200, height=18, borderStyle="inset"
    )
    c.save()
    return buf.getvalue()


def test_extract_acroform_pdf_detects_fields():
    pytest.importorskip("pypdf")
    ex = dr.extract(_acroform_pdf(), "pdf")
    assert ex.pdf_form_type == "acroform"
    assert {f.name for f in ex.fields} == {"customer_name"}
    assert all(f.source == "acroform" for f in ex.fields)


def test_render_acroform_pdf_fills_field_preserved():
    pypdf = pytest.importorskip("pypdf")
    body = _acroform_pdf()
    out, fidelity = dr.render(
        kind="pdf",
        pdf_form_type="acroform",
        body=body,
        field_values={"customer_name": "ACME Fleet"},
        replacements=[],
    )
    assert fidelity == "preserved"
    fields = pypdf.PdfReader(io.BytesIO(out)).get_fields()
    assert fields["customer_name"].get("/V") == "ACME Fleet"


def test_render_acroform_pdf_applies_replacements():
    pypdf = pytest.importorskip("pypdf")
    body = _acroform_pdf()
    out, fidelity = dr.render(
        kind="pdf",
        pdf_form_type="acroform",
        body=body,
        field_values={"customer_name": "ACME Fleet"},
        replacements=[{"find": "Customer:", "replace": "Client:"}],
    )
    assert fidelity == "preserved"
    reader = pypdf.PdfReader(io.BytesIO(out))
    assert "Client:" in (reader.pages[0].extract_text() or "")
    assert reader.get_fields()["customer_name"].get("/V") == "ACME Fleet"


def test_apply_pdf_replacements_handles_tj_array():
    # Regression: text shown via a kerned `TJ` array (typical of Word/LibreOffice
    # exports) must be replaced AND re-serialize — previously the operand was
    # wrapped in a plain list and pypdf raised AttributeError on write.
    pypdf = pytest.importorskip("pypdf")
    from pypdf.generic import DecodedStreamObject, NameObject

    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=300, height=200)
    page = writer.pages[0]
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 24 Tf 50 100 Td [(Hel) -50 (lo) ( World)] TJ ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    buf = io.BytesIO()
    writer.write(buf)

    out = dr._apply_pdf_replacements(buf.getvalue(), [{"find": "World", "replace": "Earth"}])
    raw = pypdf.PdfReader(io.BytesIO(out)).pages[0].get_contents().get_data()
    assert b"Earth" in raw and b"World" not in raw
